"""지출결의 업무 (Streamlit)

지출결의서 내역을 입력하면 아래 결과물을 자동 생성한다.
  - 품의서 하단 문구 (1. 주요내용 / 2. 사유 / 3. 첨부 / 4. 특이사항)
  - 그룹웨어 문서 제목
  - 완료기안 엑셀 (YYMMDD_완료기안 결제요청.xlsx)
  - 문자 보고용 템플릿

실행: streamlit run app.py
"""

from __future__ import annotations

import hashlib
import html
import json
import math
import os
import re
import zipfile
from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# 상수
# ---------------------------------------------------------------------------
ITEM_COLUMNS = ["품목", "업체명", "수량", "단가", "배송비", "비고"]  # 사용자가 입력하는 칸
ITEM_COMPUTED = ["순번", "공급가액", "부가세", "합계"]  # 자동 계산 칸 (편집 불가)
EDITOR_COLUMNS = ["선택", "순번", "품목", "업체명", "수량", "단가", "공급가액", "부가세", "배송비", "합계", "비고"]
EDITOR_NUMBER_COLUMNS = ["순번", "수량", "단가", "공급가액", "부가세", "배송비", "합계"]

VAT_MODES = {
    "별도 (단가 = 공급가액)": "exclusive",
    "포함 (단가 = VAT 포함가)": "inclusive",
    "면세 / 영세": "exempt",
}

# ---- 완료기안 엑셀: 원본 '*완료기안*.xlsx'의 [입금요청] 시트에 값만 채움 ----
DONE_TEMPLATE_NAME = "완료기안 결제요청.xlsx"  # 없으면 폴더의 '*완료기안*.xlsx' 중 원본(가장 큰 파일)
DONE_SHEET = "입금요청"
DONE_FIRST_ROW = 5  # 데이터 시작 행 (합계 행은 '합 계' 문구로 찾음)
DONE_ROW_HEIGHT = 39.95
DONE_LAST_COL = "AF"
DONE_STYLE_COLS = [get_column_letter(c) for c in range(2, 31)]  # B~AD
DONE_ROW_MERGES = [("D", "F"), ("G", "I"), ("J", "M"), ("N", "P"), ("R", "U"), ("Z", "AD")]
DONE_TEMPLATE_CHECK = {"B4": "순번", "C4": "문서번호", "D4": "지급처", "J4": "공급가액", "R4": "합 계",
                       "W4": "계좌번호", "X4": "입금요청일", "AA3": "GSI", "AC3": "KN"}

# ---- 지출결의서: '260821 지출결의서.xlsx'의 [2026 지출결의서] 시트 양식 ----
APP_DIR = Path(__file__).resolve().parent
EXP_TEMPLATE_NAME = "지출결의서.xlsx"  # 없으면 폴더의 '*지출결의서*.xlsx' 사용
EXP_ITEM_FIRST, EXP_ITEM_LAST = 11, 61  # 원본 내역 행
EXP_VAT_ROW, EXP_TOTAL_ROW = 62, 63
EXP_TEMPLATE_CHECK = {  # 템플릿 구조 확인용 고정 문구
    "B3": "지 출 결 의 서", "B10": "순번", "C10": "품목(사유)", "B62": "부 가 세", "B63": "합 계",
    "B64": "결제방법", "B65": "계좌정보", "B66": "법인구분", "B68": "상기와 같이 지출결의서를 작성 품의합니다.",
}
EXP_TBD = "추후별도기재"
EXP_WIDTHS = {
    "A": 2.125, "B": 5.625, "C": 9.0, "D": 9.0, "E": 9.0, "F": 21.125, "G": 6.5, "H": 11.5,
    "I": 10.625, "J": 10.5, "K": 5.625, "L": 15.625, "M": 2.125, "N": 2.125,
    "O": 10.125, "P": 9.0, "Q": 9.0, "R": 9.0, "S": 10.5,
}
EXP_GRAY = "D9D9D9"        # 순번·결제요청일·금액 (흰색 15% 어둡게)
EXP_LIGHT = "E8E8E8"       # 부가세·합계 금액
EXP_WHITE = "FFFFFF"
EXP_TITLE_FONT = "0000CC"  # '제목 :' 글자색
EXP_DATE_FONT = "215F9A"   # 결제요청일 글자색
EXP_SECTION_NAMES = {"주요내용": "주요내용", "사유": "사    유", "첨부": "첨     부", "특이사항": "특이사항"}  # 양식 띄어쓰기
EMPTY_MARKERS = {"", "없음", "-", "--", "해당없음", "해당 없음", "n/a", "na", "x", "無"}  # 이 값뿐이면 항목 생략
KO_ENUM = "가나다라마바사아자차카타파하"
PAY_METHODS = ["송금", "법인카드", "자동이체", "현금"]
EVIDENCE_TYPES = ["-", "세금계산서 - 전자", "세금계산서 - 종이", "계산서", "카드전표", "현금영수증", "영수증", "기타"]
BANKS = ["KB국민은행", "신한은행", "우리은행", "하나은행", "NH농협은행", "IBK기업은행", "SC제일은행", "한국씨티은행",
         "KDB산업은행", "Sh수협은행", "iM뱅크(대구은행)", "부산은행", "경남은행", "광주은행", "전북은행", "제주은행",
         "카카오뱅크", "케이뱅크", "토스뱅크"]  # 제1금융권 — 목록에 없으면 직접 입력

# ---- 입력 폼 (위젯 key → 초기값). 임시저장·불러오기·새로 작성이 이 key들을 사용 ----
FORM_DATE_KEYS = ("f_write_date", "f_pay_date")


def form_defaults() -> dict:
    today = date.today()
    return {
        "f_write_date": today, "f_pay_date": today, "f_payer": "GSI", "f_urgent": False,
        "f_pay_method": PAY_METHODS[0], "f_evidence": EVIDENCE_TYPES[0], "f_vat": list(VAT_MODES)[0],
        "f_bank": None, "f_account": "", "f_holder": "",
        "f_reason": "", "f_attachment": "", "f_remark": "",
    }


# ---- 임시저장 ----
DRAFTS_PATH = Path(os.environ.get("JICHUL_DRAFTS_PATH") or APP_DIR / "drafts.json")

# ---- 탭 ----
TAB_WRITE, TAB_LEDGER, TAB_DRAFTS, TAB_BACKUP = "1. 지출결의서 작성", "2. 발급 대장", "3. 임시저장 목록", "4. 백업"

# ---- 발급 대장 ----
HISTORY_PATH = Path(os.environ.get("JICHUL_HISTORY_PATH") or APP_DIR / "history.csv")  # 환경변수로 위치 변경 가능
HISTORY_COLUMNS = ["발급ID", "작성일", "문서번호", "기안부서", "기안자", "지급처", "건명", "품목수",
                   "공급가액", "부가세", "합계", "은행", "계좌번호", "예금주", "출금회사", "입금요청일", "긴급",
                   "지출결의서 발급", "완료기안 발급", "승인완료", "승인일시"]
HISTORY_NUMBER_COLUMNS = ["품목수", "공급가액", "부가세", "합계"]
STATUS_DRAFT = "🟡 1. 기안작성 중"          # 문서번호 없음
STATUS_SUBMITTED = "🔵 2. 상신완료/진행 중"  # 문서번호 입력 또는 완료기안 발급
STATUS_DONE = "🟢 3. 최종승인 완료"          # 승인완료 체크
STATUS_FILTERS = {"전체": None, "기안작성": STATUS_DRAFT, "상신완료": STATUS_SUBMITTED, "최종승인완료": STATUS_DONE}
STATUS_LEGEND = "🟡 기안 작성 중 (문서번호 미입력) · 🔵 상신 완료/결재 진행 중 (문서번호 입력) · 🟢 최종 승인 완료"
LEDGER_VIEW_COLUMNS = ["상태", "건명", "작성일", "문서번호", "승인완료", "지급처", "합계", "공급가액", "부가세",
                       "품목수", "기안자", "출금회사", "입금요청일", "긴급", "은행", "계좌번호", "예금주", "기안부서",
                       "지출결의서 발급", "완료기안 발급", "승인일시"]

# ---- 백업: history.csv + drafts.json을 zip으로 묶어 backups 폴더에 보관 ----
BACKUP_DIR = Path(os.environ.get("JICHUL_BACKUP_DIR") or APP_DIR / "backups")
BACKUP_KEEP = 60  # 자동 백업은 최근 60개만 보관 (수동·복원/삭제 직전 백업은 지우지 않음)
BACKUP_KINDS = {"auto": "자동(하루 1회)", "manual": "수동", "before-restore": "복원 직전", "before-delete": "삭제 직전"}


# ---------------------------------------------------------------------------
# 계산
# ---------------------------------------------------------------------------
def _num(value) -> float:
    try:
        v = float(value)
    except (TypeError, ValueError):
        return 0.0
    return 0.0 if math.isnan(v) else v


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def _opt_num(value) -> float:
    """빈칸은 NaN 그대로 (표에서 빈칸으로 보이도록)."""
    try:
        v = float(value)
    except (TypeError, ValueError):
        return math.nan
    return v


def _flag(value) -> bool:
    return value is not None and not (isinstance(value, float) and math.isnan(value)) and bool(value)


def _amounts(qty: float, price: float, ship: int, vat_mode: str) -> tuple[int, int]:
    """(공급가액, 부가세). 양식 기준: 배송비도 과세 대상 → 공급가액 = 수량×단가 + 배송비, 부가세 = 공급가액의 10%
    (예: 지출결의서 소계 418,000 → 부가세 41,800 / 완료기안 공급가액 418,000)."""
    gross = int(round(qty * price)) + ship
    if vat_mode == "exclusive":
        return gross, int(gross * 0.1)
    if vat_mode == "inclusive":
        supply = int(round(gross / 1.1))
        return supply, gross - supply
    return gross, 0


def blank_items(n: int = 1) -> pd.DataFrame:
    return normalize_items(pd.DataFrame([{} for _ in range(n)], columns=EDITOR_COLUMNS), "exclusive")


def normalize_items(df: pd.DataFrame, vat_mode: str) -> pd.DataFrame:
    """내역 입력 표 정리 + 자동 계산 칸(순번·공급가액·부가세·합계) 채움. 같은 입력이면 항상 같은 결과."""
    rows, seq = [], 0
    for r in df.to_dict("records"):
        name = _text(r.get("품목"))
        row = {
            "선택": _flag(r.get("선택")), "순번": math.nan, "품목": name, "업체명": _text(r.get("업체명")),
            "수량": _opt_num(r.get("수량")), "단가": _opt_num(r.get("단가")),
            "공급가액": math.nan, "부가세": math.nan, "배송비": _opt_num(r.get("배송비")), "합계": math.nan,
            "비고": _text(r.get("비고")),
        }
        if name:
            seq += 1
            supply, vat = _amounts(_num(row["수량"]), _num(row["단가"]), int(round(_num(row["배송비"]))), vat_mode)
            row.update({"순번": seq, "공급가액": supply, "부가세": vat, "합계": supply + vat})
        rows.append(row)
    out = pd.DataFrame(rows, columns=EDITOR_COLUMNS)
    out["선택"] = out["선택"].astype(bool)
    out[EDITOR_NUMBER_COLUMNS] = out[EDITOR_NUMBER_COLUMNS].astype(float)
    for col in ("품목", "업체명", "비고"):
        out[col] = out[col].astype(str)
    return out


def compute_items(df: pd.DataFrame, vat_mode: str, default_vendor: str = "") -> pd.DataFrame:
    """입력 표 → 공급가액/부가세/배송비/합계가 계산된 표 (빈 행 제외)."""
    rows = []
    for _, r in df.iterrows():
        name = _text(r.get("품목"))
        if not name:
            continue
        qty = _num(r.get("수량"))
        price = _num(r.get("단가"))
        ship = int(round(_num(r.get("배송비"))))
        supply, vat = _amounts(qty, price, ship, vat_mode)

        rows.append({
            "순번": len(rows) + 1,
            "품목": name,
            "업체명": _text(r.get("업체명")) or default_vendor,
            "수량": qty,
            "단가": price,
            "공급가액": supply,
            "부가세": vat,
            "배송비": ship,
            "합계": supply + vat,
            "비고": _text(r.get("비고")),
        })
    return pd.DataFrame(rows, columns=[
        "순번", "품목", "업체명", "수량", "단가", "공급가액", "부가세", "배송비", "합계", "비고",
    ])


def won(n: float) -> str:
    return f"{int(n):,}원"


def qty_str(q: float) -> str:
    return f"{int(q):,}" if float(q).is_integer() else f"{q:,}"


def item_summary(items: pd.DataFrame) -> str:
    """건명 = 내역 첫 번째 행의 품목 (내역이 없으면 빈 문자열)."""
    if items.empty or "품목" not in items.columns:
        return ""
    return _text(items.iloc[0]["품목"])


def with_gun(subject: str) -> str:
    """'명함 제작' → '명함 제작 건', 이미 '건'으로 끝나면 그대로."""
    if not subject:
        return ""
    return subject if subject.endswith("건") else f"{subject} 건"


def has_content(text: str) -> bool:
    """'없음', '-' 등 내용이 없다는 표시만 있으면 False."""
    lines = [ln.strip().rstrip(".").strip() for ln in (text or "").splitlines()]
    return any(ln and ln.lower() not in EMPTY_MARKERS for ln in lines)


def _section_lines(text: str) -> list[str]:
    """입력란 내용 → '  가. …' 줄들 (이미 번호가 있으면 그대로)."""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    return ["  " + ln for ln in _enumerate_ko(lines)]


# ---------------------------------------------------------------------------
# 문구 생성
# ---------------------------------------------------------------------------
def build_body_text(info: dict, items: pd.DataFrame, totals: dict) -> str:
    """품의 문구. 사유·첨부·특이사항은 내용이 있을 때만 넣고 번호를 1부터 다시 매김."""
    subject = info["subject"]

    if items.empty:
        qty_line = "-"
    elif len(items) == 1:
        qty_line = f"{qty_str(items.iloc[0]['수량'])}{info['unit']}"
    else:
        parts = [f"{r['품목']} {qty_str(r['수량'])}{info['unit']}" for _, r in items.iterrows()]
        qty_line = f"총 {len(items)}건 (" + ", ".join(parts) + ")"

    vat_note = {"exclusive": "VAT 포함", "inclusive": "VAT 포함", "exempt": "면세"}[info["vat_mode"]]
    notes = [vat_note]
    if totals["배송비"]:
        notes.append(f"배송비 {won(totals['배송비'])} 포함")
    amount_line = f"￦ {won(totals['합계'])} ({', '.join(notes)})"

    account_line = " ".join(x for x in [info["bank"], info["account"]] if x)
    if info["holder"]:
        account_line += f" (예금주: {info['holder']})"

    main = [
        f"  가. 품목: {with_gun(subject)}",
        f"  나. 수량: {qty_line}",
        f"  다. 금액: {amount_line}",
        f"  라. 계좌번호: {account_line or EXP_TBD}",
    ]
    if len(items) > 1:
        main.append("")
        main.append("  ※ 세부내역")
        for _, r in items.iterrows():
            main.append(
                f"    {r['순번']}) {r['품목']} / {r['업체명']} / "
                f"{qty_str(r['수량'])}{info['unit']} / {won(r['합계'])}"
            )

    sections = [("주요내용", main)]
    for name, key in (("사유", "reason"), ("첨부", "attachment"), ("특이사항", "remark")):
        if has_content(info[key]):
            sections.append((name, _section_lines(info[key])))

    blocks = [f"{i}. {name}\n" + "\n".join(lines) for i, (name, lines) in enumerate(sections, start=1)]
    return "\n\n".join(blocks) + "\n끝."


def build_title(info: dict) -> str:
    tags = "".join(f"[{t}]" for t in [info["company"], info["site"], info["dept"]] if t)
    return f"{tags} {with_gun(info['subject'])}."


def build_sms(info: dict, items: pd.DataFrame) -> str:
    """문자 보고. 내역 행마다 번호 · 업체명 · 품목 · 비용(그 행의 합계, VAT 포함)."""
    blocks = [
        f"{i}. {r['업체명'] or '(업체명 미입력)'}\n- 품목: {r['품목']}\n- 비용: {won(r['합계'])}"
        for i, r in enumerate(items.to_dict("records"), start=1)
    ]
    return (
        f"{info['sms_to']}, 안녕하세요.\n"
        f"{info['sms_from']}입니다.\n"
        "\n"
        "결제승인 된 기안 입금 요청건으로 보고드립니다.\n"
        "\n"
        + "\n\n".join(blocks)
        + "\n\n감사합니다."
    )


# ---------------------------------------------------------------------------
# 엑셀 생성
# ---------------------------------------------------------------------------
def _enumerate_ko(lines: list[str]) -> list[str]:
    """'가.' '나.' … 번호가 없는 줄에만 번호를 붙임."""
    out, k = [], 0
    for line in lines:
        if re.match(r"^([가나다라마바사아자차카타파하]\.|\d+[.)]|[①-⑳]|※|-\s|\*)", line):
            out.append(line)
        else:
            out.append(f"{KO_ENUM[k % len(KO_ENUM)]}. {line}")
            k += 1
    return out


def parse_body_sections(body: str) -> list[tuple[str, str]]:
    """품의 문구 → [(제목행, 본문)] — 줄 맨 앞의 'N. ' 기준으로 나눔. 사용자가 수정한 내용 그대로 반영."""
    sections: list[tuple[str, str, list[str]]] = []
    lines = [re.sub(r"\s*끝\.\s*$", "", ln) for ln in body.replace("\r\n", "\n").split("\n")]
    for line in lines:
        m = re.match(r"^(\d+)\.\s*([^:：]*)[:：]?\s*(.*)$", line)
        if m and not line.startswith(" "):
            num, name, inline = m.group(1), m.group(2).strip(), m.group(3).strip()
            header = f"{num}. {EXP_SECTION_NAMES.get(name.replace(' ', ''), name)}"
            sections.append((header, inline, []))
        elif sections:
            sections[-1][2].append(line)
        elif line.strip():
            sections.append(("", line.strip(), []))

    result = []
    for header, inline, rest in sections:
        if inline:  # '2. 사유: …' 형식 → 본문을 '가. …'로
            lines = [inline] + [l.strip() for l in rest]
            lines = [l for l in lines if l]
            lines = [re.sub(r"\s*끝\.$", "", l) for l in lines]
            lines = [l for l in lines if l]
            text = "\n".join("  " + l for l in _enumerate_ko(lines))
        else:  # '1. 주요내용' + 들여쓴 줄들 → 그대로
            while rest and not rest[-1].strip():
                rest.pop()
            while rest and not rest[0].strip():
                rest.pop(0)
            text = "\n".join(re.sub(r"\s*끝\.$", "", l) for l in rest)
        result.append((header, text))
    return result


def find_expense_template() -> Path | None:
    """app.py 폴더에서 지출결의서 원본 템플릿 찾기: '지출결의서.xlsx' 우선, 없으면 '*지출결의서*.xlsx'.
    앱이 만든 다운로드 파일(YYMMDD_지출결의서.xlsx)은 후순위."""
    exact = APP_DIR / EXP_TEMPLATE_NAME
    if exact.exists():
        return exact
    found = [p for p in APP_DIR.glob("*지출결의서*.xlsx") if not p.name.startswith("~$")]
    found.sort(key=lambda p: (bool(re.match(r"^\d{6}_지출결의서\.xlsx$", p.name)), p.name))
    return found[0] if found else None


def check_expense_template(path: Path) -> str | None:
    """템플릿 구조 확인. 문제가 있으면 오류 문구, 정상이면 None."""
    try:
        ws = load_workbook(path).worksheets[0]
    except Exception as exc:  # noqa: BLE001 — 사용자에게 원인 그대로 표시
        return f"템플릿을 열 수 없습니다: {exc}"
    for coord, expected in EXP_TEMPLATE_CHECK.items():
        if str(ws[coord].value or "").replace(" ", "") != expected.replace(" ", ""):
            return f"템플릿 구조가 다릅니다: {coord} 셀이 '{expected}'이어야 합니다 (현재: {ws[coord].value!r})."
    return None


def _copy_row_style(ws, src_row: int, dst_row: int, cols: str = "BCDEFGHIJKL") -> None:
    """원본 템플릿의 행 서식을 그대로 복사 (새 서식을 만들지 않음)."""
    for col in cols:
        ws[f"{col}{dst_row}"]._style = copy(ws[f"{col}{src_row}"]._style)


def _text_height(text: str, width: float, line_pt: float, minimum: float) -> float:
    """병합 셀은 Excel이 높이를 자동 맞춤하지 않으므로 줄 수로 높이 계산."""
    lines = sum(max(1, math.ceil(sum(2 if ord(ch) > 127 else 1 for ch in ln) / max(width - 1, 1)))
                for ln in str(text).split("\n"))
    return max(minimum, line_pt * lines + 4)


def build_expense_excel(template: Path, info: dict, items: pd.DataFrame, totals: dict,
                        body: str, title: str) -> bytes:
    """원본 지출결의서 템플릿을 불러와 지정 셀에 값만 채움 (서식·병합·열 너비·인쇄 설정은 원본 그대로).

    - 내역: 11행부터 품목 수만큼 작성, 남는 내역 행(최대 61행까지)은 숨김 → 원본 구조 보존
    - 품의 문구: 1~3번은 원본 70~75행, 4. 특이사항은 원본의 빈 76~77행에 같은 서식으로 이어서 작성
    """
    wb = load_workbook(template)
    ws = wb.worksheets[0]
    first, last = EXP_ITEM_FIRST, EXP_ITEM_LAST
    # 열 너비 읽기 (column_dimensions[키] 접근은 묶음 열에 새 항목을 만들어 너비를 바꾸므로 사용 금지)
    col_w = {get_column_letter(c): d.width for d in list(ws.column_dimensions.values())
             for c in range(d.min, d.max + 1)}
    width = lambda cols: sum(col_w.get(c) or 9.0 for c in cols)  # noqa: E731

    # ---- 제목 / 기본 정보 ----
    ws["B5"] = f"제목 : {title}"
    ws["D6"] = info["write_date"]
    ws["G6"] = info["dept"]
    ws["J6"] = info["site"]
    ws["D7"] = info["pay_date"]
    ws["G7"] = info["drafter"]
    ws["J7"] = info["requester"]
    ws["I8"] = f"=I{EXP_TOTAL_ROW}"

    # ---- 내역 (11행~) — 서식은 원본 11행(입력 행) 서식을 복사 ----
    records = items.to_dict("records")
    base_height = ws.row_dimensions[first].height or 45.0
    for i, r in enumerate(range(first, last + 1)):
        rec = records[i] if i < len(records) else None
        _copy_row_style(ws, first, r)
        ws[f"B{r}"] = f'=IF(C{r}="","",ROW()-{first - 1})'
        for col in "CFGHIJK":
            ws[f"{col}{r}"] = None
        if rec is None:  # 남는 행: 비우고 숨김
            ws.row_dimensions[r].hidden = True
            continue
        ws.row_dimensions[r].hidden = False
        ws[f"C{r}"] = rec["품목"]
        ws[f"F{r}"] = rec["업체명"]
        ws[f"G{r}"] = int(rec["수량"]) if float(rec["수량"]).is_integer() else rec["수량"]
        ws[f"H{r}"] = int(rec["단가"]) if float(rec["단가"]).is_integer() else rec["단가"]
        ws[f"I{r}"] = int(rec["배송비"])
        ws[f"J{r}"] = f"=(G{r}*H{r})+I{r}"
        ws[f"K{r}"] = rec["비고"] or "-"
        ws.row_dimensions[r].height = max(
            base_height,
            _text_height(rec["품목"], width("CDE"), 15.0, 0),
            _text_height(rec["업체명"], width("F"), 15.0, 0),
            _text_height(rec["비고"], width("KL"), 15.0, 0),
        )

    # ---- 부가세 / 합계 ----
    inclusive = info["vat_mode"] == "inclusive"  # 단가에 VAT 포함 → 소계에 이미 들어 있음
    ws[f"I{EXP_VAT_ROW}"] = 0 if inclusive else int(totals["부가세"])
    ws[f"K{EXP_VAT_ROW}"] = "단가에 VAT 포함" if inclusive else "-"
    ws[f"I{EXP_TOTAL_ROW}"] = f"=SUM(J{first}:J{last},I{EXP_VAT_ROW})"

    # ---- 결제방법 / 계좌정보 / 법인구분 ----
    ws["D64"] = info["pay_method"]
    ws["G64"] = info["evidence"]
    ws["E65"] = info["bank"] or "-"
    ws["G65"] = info["account"] or EXP_TBD
    ws["K65"] = info["holder"] or EXP_TBD
    ws["D66"] = f"  {info['payer']}  입금 요청"

    # ---- 품의 문구 (1~4번 + 끝.) ----
    # 원본: 70/71(1번) · 72/73(2번) · 74/75(3번) · 76(끝.) — 4번은 76/77에 원본 제목·본문 서식 복사, 끝.은 78행(원본 병합 B78:L78)
    sections = parse_body_sections(body)
    if len(sections) > 4:  # 5번 이상을 추가했다면 4번 칸에 이어 붙임
        extra = "\n".join(f"{h}\n{t}" if h else t for h, t in sections[3:])
        sections = sections[:3] + [("", extra)]
    end_style = copy(ws["B76"]._style)  # 원본 '끝.' 서식 (덮어쓰기 전에 보관)
    if "B77:L77" not in {str(m) for m in ws.merged_cells.ranges}:
        ws.merge_cells("B77:L77")
    _copy_row_style(ws, 74, 76, "B")  # 4번 제목 ← 3번 제목 서식
    _copy_row_style(ws, 75, 77, "B")  # 4번 본문 ← 3번 본문 서식

    body_width = width("BCDEFGHIJKL")
    header_height = ws.row_dimensions[70].height or 17.25
    for i in range(4):
        h_row, b_row = 70 + 2 * i, 71 + 2 * i
        header, text = sections[i] if i < len(sections) else ("", "")
        ws[f"B{h_row}"] = header or None
        ws[f"B{b_row}"] = text or None
        ws.row_dimensions[h_row].height = header_height
        ws.row_dimensions[b_row].height = _text_height(text, body_width, 16.5, 36.0) if text else None
        hide = i >= len(sections)
        ws.row_dimensions[h_row].hidden = hide
        ws.row_dimensions[b_row].hidden = hide
    ws["B78"]._style = end_style
    ws["B78"] = "끝."
    ws.row_dimensions[78].hidden = False

    # ---- 인쇄 영역: 원본 A1:M76 → 끝.(78행)까지 ----
    ws.print_area = "A1:M78"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def find_done_template() -> Path | None:
    """app.py 폴더에서 완료기안 원본 찾기: '완료기안 결제요청.xlsx' 우선, 없으면 '*완료기안*.xlsx' 중 가장 큰 파일
    (원본은 여러 날짜 시트가 쌓인 파일이라 앱이 만든 한 장짜리 다운로드 파일보다 큼)."""
    exact = APP_DIR / DONE_TEMPLATE_NAME
    if exact.exists():
        return exact
    found = [p for p in APP_DIR.glob("*완료기안*.xlsx") if not p.name.startswith("~$")]
    return max(found, key=lambda p: p.stat().st_size) if found else None


@st.cache_data(show_spinner="완료기안 템플릿 불러오는 중…")
def load_done_template(path: str, mtime: float) -> bytes:
    """원본에서 [입금요청] 시트만 남긴 템플릿 (원본 파일은 수정하지 않음).
    원본에 시트가 수십 개라 열기가 느리므로 파일 수정 시각(mtime) 기준으로 캐시."""
    wb = load_workbook(path)
    if DONE_SHEET not in wb.sheetnames:
        raise ValueError(f"'{DONE_SHEET}' 시트가 없습니다 (시트: {', '.join(wb.sheetnames[:5])} …).")
    for name in list(wb.sheetnames):
        if name != DONE_SHEET:
            del wb[name]
    ws = wb[DONE_SHEET]
    ws.sheet_state = "visible"
    wb.active = 0
    # 원본의 첫 탭 위치(firstSheet=17 등)가 남아 있으면 시트 1장짜리 파일에서 Excel 인쇄/PDF가 오류 → 초기화
    for view in wb.views:
        view.firstSheet = 0
        view.activeTab = 0
    for coord, expected in DONE_TEMPLATE_CHECK.items():
        if str(ws[coord].value or "").replace(" ", "") != expected.replace(" ", ""):
            raise ValueError(f"템플릿 구조가 다릅니다: {coord} 셀이 '{expected}'이어야 합니다 (현재: {ws[coord].value!r}).")
    _done_total_row(ws)  # 합계 행이 있는지 확인
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _done_total_row(ws) -> int:
    for r in range(DONE_FIRST_ROW, ws.max_row + 1):
        if str(ws[f"B{r}"].value or "").replace(" ", "") == "합계":
            return r
    raise ValueError("템플릿에서 '합 계' 행을 찾을 수 없습니다.")


def build_excel(template: bytes, info: dict, items: pd.DataFrame, totals: dict) -> bytes:
    """완료기안 원본 [입금요청] 시트에 값만 채움 (서식·병합·열 너비·인쇄 설정은 원본 그대로).

    - 데이터: 5행부터 품목 수만큼. 원본 데이터 칸(5행)보다 많으면 합계 행을 아래로 옮기고 원본 5행 서식으로 행 추가
    - 남는 데이터 행은 비우고 숨김
    """
    wb = load_workbook(BytesIO(template))
    ws = wb[DONE_SHEET]
    first = DONE_FIRST_ROW
    total_row = _done_total_row(ws)
    records = items.to_dict("records")
    extra = len(records) - (total_row - first)

    # ---- 행이 모자라면: 합계 행을 아래로 이동 + 원본 5행 서식으로 데이터 행 추가 ----
    if extra > 0:
        total_merges = [m for m in list(ws.merged_cells.ranges) if m.min_row == total_row]
        total_height = ws.row_dimensions[total_row].height
        for m in total_merges:
            ws.unmerge_cells(str(m))
        ws.move_range(f"A{total_row}:{DONE_LAST_COL}{total_row}", rows=extra)
        for m in total_merges:
            ws.merge_cells(start_row=total_row + extra, start_column=m.min_col,
                           end_row=total_row + extra, end_column=m.max_col)
        ws.row_dimensions[total_row + extra].height = total_height
        for r in range(total_row, total_row + extra):
            _copy_row_style(ws, first, r, DONE_STYLE_COLS)
            for start, end in DONE_ROW_MERGES:
                ws.merge_cells(f"{start}{r}:{end}{r}")
        total_row += extra

    # ---- 출금회사 범례: 선택된 회사 칸에 'o' (원본 Z3=GSI, AB3=KN) ----
    is_gsi = info["payer"] == "GSI"
    ws["Z3"] = "o" if is_gsi else None
    ws["AB3"] = None if is_gsi else "o"
    payer_fill = copy(ws["Z3" if is_gsi else "AB3"].fill)  # 순번 칸 배경 = 범례 색

    # ---- 데이터 행 ----
    col_w = {get_column_letter(c): d.width for d in list(ws.column_dimensions.values())
             for c in range(d.min, d.max + 1)}
    width = lambda cols: sum(col_w.get(c) or 9.0 for c in cols)  # noqa: E731
    for i, r in enumerate(range(first, total_row)):
        rec = records[i] if i < len(records) else None
        _copy_row_style(ws, first, r, DONE_STYLE_COLS)
        for col in ("C", "D", "G", "J", "N", "Q", "V", "W", "X", "Y", "Z"):
            ws[f"{col}{r}"] = None
        ws[f"B{r}"] = i + 1
        ws[f"R{r}"] = f"=SUM(J{r}:Q{r})"
        if rec is None:
            ws.row_dimensions[r].hidden = True
            continue
        ws.row_dimensions[r].hidden = False
        ws[f"B{r}"].fill = copy(payer_fill)
        ws[f"C{r}"] = info["doc_no"] or None
        ws[f"D{r}"] = rec["업체명"]
        ws[f"G{r}"] = rec["품목"]
        ws[f"J{r}"] = int(rec["공급가액"])  # 배송비 포함 → 배송비(Q) 칸은 원본처럼 비움
        ws[f"N{r}"] = int(rec["부가세"])
        ws[f"V{r}"] = info["bank"] or None
        ws[f"W{r}"] = info["account"] or None
        ws[f"X{r}"] = info["pay_date"]
        ws[f"Y{r}"] = "o" if info["urgent"] else "x"
        ws[f"Z{r}"] = rec["비고"] or info["payer"]
        ws.row_dimensions[r].height = max(
            DONE_ROW_HEIGHT,
            _text_height(rec["업체명"], width("DEF"), 18.0, 0),
            _text_height(rec["품목"], width("GHI"), 18.0, 0),
            _text_height(rec["비고"], width(["Z", "AA", "AB", "AC", "AD"]), 18.0, 0),
        )

    # ---- 합계 / 인쇄 영역 ----
    ws[f"R{total_row}"] = f"=SUM(R{first}:U{total_row - 1})"
    ws.print_area = f"B2:AD{total_row}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 지출결의서 HTML 미리보기 (그룹웨어 붙여넣기용 — 모든 서식을 인라인 style로)
# ---------------------------------------------------------------------------
def korean_amount(n: int) -> str:
    """Excel NUMBERSTRING(n, 1)과 같은 한글 금액 (예: 459800 → 사십오만구천팔백)."""
    if n <= 0:
        return "영"
    digits = "영일이삼사오육칠팔구"
    small = ["", "십", "백", "천"]
    big = ["", "만", "억", "조", "경"]
    out, group = [], 0
    while n > 0:
        n, chunk = divmod(n, 10000)
        if chunk:
            part, pos = "", 0
            while chunk > 0:
                chunk, d = divmod(chunk, 10)
                if d:
                    part = digits[d] + small[pos] + part
                pos += 1
            out.insert(0, part + big[group])
        group += 1
    return "".join(out)


def _h(text) -> str:
    """HTML 이스케이프 + 줄바꿈/앞 공백 보존 (편집기가 white-space를 지워도 유지되도록)."""
    if text is None:
        return ""
    lines = []
    for line in str(text).split("\n"):
        stripped = line.lstrip(" ")
        escaped = html.escape(stripped).replace("  ", "&nbsp; ")  # '2. 사    유' 같은 연속 공백 유지
        lines.append("&nbsp;" * (len(line) - len(stripped)) + escaped)
    return "<br>".join(lines)


def build_expense_html(info: dict, items: pd.DataFrame, totals: dict, body: str, title: str) -> str:
    """지출결의서 엑셀 양식과 같은 구성의 HTML 표 (결재란 · 기본정보 · 내역 · 품의 문구)."""
    font = "font-family:'맑은 고딕','Malgun Gothic',sans-serif;"
    bd = "border:1px solid #000;"
    base = f"{font}{bd}font-size:10pt;padding:4px 6px;text-align:center;vertical-align:middle;word-break:keep-all;"

    def td(content="", *, colspan=1, rowspan=1, bold=False, bg=None, color=None, align=None,
           size=None, height=None, extra="") -> str:
        style = base
        if bold:
            style += "font-weight:bold;"
        if bg:
            style += f"background-color:#{bg};"
        if color:
            style += f"color:#{color};"
        if align:
            style += f"text-align:{align};"
        if size:
            style += f"font-size:{size}pt;"
        if height:
            style += f"height:{height}px;"
        attrs = (f' colspan="{colspan}"' if colspan > 1 else "") + (f' rowspan="{rowspan}"' if rowspan > 1 else "")
        return f'<td{attrs} style="{style}{extra}">{content}</td>'

    def num(v) -> str:
        return f"{int(v):,}" if v else "-"

    cols = ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]
    total_w = sum(EXP_WIDTHS[c] for c in cols)
    colgroup = "".join(f'<col style="width:{EXP_WIDTHS[c] / total_w * 100:.2f}%">' for c in cols)
    inclusive = info["vat_mode"] == "inclusive"
    vat_shown = 0 if inclusive else int(totals["부가세"])

    rows = [
        # 제목 / 결재란
        "<tr>" + td("지 출 결 의 서", colspan=9, rowspan=2, bold=True, size=24)
        + td("결<br>재", rowspan=3, bold=True) + td("대표이사", bold=True, height=34) + "</tr>",
        "<tr>" + td("", rowspan=2, height=60) + "</tr>",
        "<tr>" + td(f"제목 : {_h(title)}", colspan=9, bold=True, color=EXP_TITLE_FONT, align="left", size=9)
        + "</tr>",
        # 기본 정보
        "<tr>" + td("작성일", colspan=2, bold=True) + td(f"{info['write_date']:%Y-%m-%d}", colspan=2)
        + td("기안부서", bold=True) + td(_h(info["dept"]), colspan=2) + td("사업소명", bold=True)
        + td(_h(info["site"]), colspan=3) + "</tr>",
        "<tr>" + td("결제요청일", colspan=2, bold=True)
        + td(f"{info['pay_date']:%Y-%m-%d}", colspan=2, bold=True, bg=EXP_GRAY, color=EXP_DATE_FONT)
        + td("기안자", bold=True) + td(_h(info["drafter"]), colspan=2) + td("요청자", bold=True)
        + td(_h(info["requester"]), colspan=3) + "</tr>",
        "<tr>" + td("금액", colspan=2, bold=True) + td("일금", bold=True)
        + td(f"{korean_amount(int(totals['합계']))}원 정", colspan=4, bold=True, bg=EXP_GRAY)
        + td(f"₩{int(totals['합계']):,}", colspan=4, bold=True, bg=EXP_GRAY) + "</tr>",
        # 내역
        "<tr>" + td("내 역", colspan=11, bold=True) + "</tr>",
        "<tr>" + td("순번", bold=True) + td("품목(사유)", colspan=3, bold=True) + td("업체명", bold=True)
        + td("수량", bold=True) + td("단가", bold=True) + td("배송비", bold=True) + td("소계", bold=True)
        + td("비고", colspan=2, bold=True) + "</tr>",
    ]
    for rec in items.to_dict("records"):
        subtotal = int(round(rec["수량"] * rec["단가"])) + int(rec["배송비"])
        rows.append(
            "<tr>" + td(rec["순번"], bg=EXP_GRAY, height=34) + td(_h(rec["품목"]), colspan=3)
            + td(_h(rec["업체명"])) + td(qty_str(rec["수량"])) + td(num(rec["단가"]))
            + td(num(rec["배송비"])) + td(num(subtotal)) + td(_h(rec["비고"] or "-"), colspan=2) + "</tr>"
        )
    rows += [
        "<tr>" + td("부 가 세", colspan=7, bold=True) + td(f"₩{vat_shown:,}", colspan=2, bold=True, bg=EXP_LIGHT)
        + td("단가에 VAT 포함" if inclusive else "-", colspan=2) + "</tr>",
        "<tr>" + td("합 계", colspan=7, bold=True)
        + td(f"₩{int(totals['합계']):,}", colspan=2, bold=True, bg=EXP_LIGHT) + td("-", colspan=2) + "</tr>",
        "<tr>" + td("결제방법", colspan=2, bold=True) + td(_h(info["pay_method"]), colspan=2)
        + td("증빙구분", bold=True) + td(_h(info["evidence"]), colspan=6) + "</tr>",
        "<tr>" + td("계좌정보", colspan=2, bold=True) + td("은행", bold=True) + td(_h(info["bank"] or "-"))
        + td("계좌번호", bold=True) + td(_h(info["account"] or EXP_TBD), colspan=3)
        + td("예금주", bold=True) + td(_h(info["holder"] or EXP_TBD), colspan=2) + "</tr>",
        "<tr>" + td("법인구분", colspan=2, bold=True)
        + td(f"&nbsp;&nbsp;{_h(info['payer'])}&nbsp;&nbsp;입금 요청", colspan=9, bold=True, align="left") + "</tr>",
    ]

    p = f"{font}margin:0;text-align:left;"
    text = [f'<p style="{p}font-size:12pt;font-weight:bold;margin-top:18px;">상기와 같이 지출결의서를 작성 품의합니다.</p>']
    for header, section in parse_body_sections(body):
        if header:
            text.append(f'<p style="{p}font-size:12pt;font-weight:bold;margin-top:12px;">{_h(header)}</p>')
        if section:
            text.append(f'<p style="{p}font-size:11pt;margin-top:4px;line-height:1.6;">{_h(section)}</p>')
    text.append(f'<p style="{p}font-size:11pt;margin-top:10px;">끝.</p>')

    return (
        f'<div style="{font}max-width:820px;margin:0 auto;color:#000;">'
        f'<p style="{p}font-size:10pt;margin-bottom:4px;">(단위 : KRW)</p>'
        f'<table style="{font}border-collapse:collapse;width:100%;table-layout:fixed;">'
        f"<colgroup>{colgroup}</colgroup><tbody>{''.join(rows)}</tbody></table>"
        f"{''.join(text)}</div>"
    )


def html_preview(doc_html: str, height: int) -> None:
    """미리보기 iframe + [HTML 복사하기] 버튼 (서식 있는 HTML과 HTML 소스를 함께 클립보드에 넣음)."""
    payload = json.dumps(doc_html)
    st.iframe(
        f"""
        <div style="position:sticky;top:0;z-index:1;background:#fff;padding:6px 0 10px;">
          <button id="copy-html" onmouseover="this.style.background='#7A895F';this.style.borderColor='#7A895F'" onmouseout="this.style.background='#8B9A6E';this.style.borderColor='#8B9A6E'" style="
              width:100%;padding:0.55rem 0.75rem;border-radius:0.5rem;cursor:pointer;
              border:1px solid #8B9A6E;background:#8B9A6E;color:#fff;font-weight:600;
              font-size:0.95rem;font-family:sans-serif;">📋 HTML 복사하기</button>
        </div>
        <div id="doc" style="background:#fff;padding:12px;border:1px solid #E0E0E0;">{doc_html}</div>
        <script>
        const html = {payload};
        const btn = document.getElementById("copy-html");
        function viaCopyEvent() {{
            // 서식(text/html) + 소스(text/plain)를 동시에 설정 → 게시판 편집기/HTML 모드 모두 붙여넣기 가능
            let ok = false;
            const onCopy = (e) => {{
                e.clipboardData.setData("text/html", html);
                e.clipboardData.setData("text/plain", html);
                e.preventDefault();
                ok = true;
            }};
            document.addEventListener("copy", onCopy);
            try {{ document.execCommand("copy"); }} catch (e) {{}}
            document.removeEventListener("copy", onCopy);
            return ok;
        }}
        async function viaClipboardApi() {{
            if (!(navigator.clipboard && window.ClipboardItem && window.isSecureContext)) return false;
            try {{
                await navigator.clipboard.write([new ClipboardItem({{
                    "text/html": new Blob([html], {{type: "text/html"}}),
                    "text/plain": new Blob([html], {{type: "text/plain"}}),
                }})]);
                return true;
            }} catch (e) {{ return false; }}
        }}
        btn.addEventListener("click", async () => {{
            const ok = viaCopyEvent() || await viaClipboardApi();
            btn.textContent = ok ? "✅ HTML 복사 완료! 그룹웨어 편집기에 붙여넣으세요" : "⚠️ 복사 실패 - 미리보기를 드래그해 복사하세요";
            setTimeout(() => btn.textContent = "📋 HTML 복사하기", 2200);
        }});
        </script>
        """,
        height=height,
    )


# ---------------------------------------------------------------------------
# 발급 대장 (history.csv 누적)
# ---------------------------------------------------------------------------
def load_history() -> pd.DataFrame:
    if not HISTORY_PATH.exists():
        return pd.DataFrame(columns=HISTORY_COLUMNS)
    df = pd.read_csv(HISTORY_PATH, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    for col in HISTORY_COLUMNS:  # 예전 파일에 없는 컬럼 보충
        if col not in df.columns:
            df[col] = ""
    df = df[HISTORY_COLUMNS]
    for col in HISTORY_NUMBER_COLUMNS:
        df[col] = pd.to_numeric(df[col].str.replace(",", ""), errors="coerce").fillna(0).astype(int)
    return df


def save_history(df: pd.DataFrame) -> None:
    """임시 파일에 쓴 뒤 교체 (저장 중 오류로 대장이 깨지지 않도록). Excel에서 한글이 보이도록 utf-8-sig."""
    tmp = HISTORY_PATH.with_suffix(".csv.tmp")
    df[HISTORY_COLUMNS].to_csv(tmp, index=False, encoding="utf-8-sig")
    tmp.replace(HISTORY_PATH)


def make_history_record(info: dict, items: pd.DataFrame, totals: dict) -> dict:
    vendors = list(dict.fromkeys(v for v in items["업체명"] if v))
    squash = lambda t: " ".join(str(t).split())  # noqa: E731 — 줄바꿈·공백 차이는 같은 건으로 봄
    key = "|".join([f"{info['write_date']:%Y-%m-%d}", squash(info["subject"]), squash(",".join(vendors)),
                    str(totals["합계"]), squash(info["drafter"])])
    return {
        "발급ID": f"{info['write_date']:%y%m%d}-{hashlib.sha1(key.encode()).hexdigest()[:6]}",
        "작성일": f"{info['write_date']:%Y-%m-%d}",
        "문서번호": info["doc_no"],
        "기안부서": info["dept"],
        "기안자": info["drafter"],
        "지급처": ", ".join(v for v in vendors if v),
        "건명": info["subject"],
        "품목수": len(items),
        "공급가액": int(totals["공급가액"]),
        "부가세": int(totals["부가세"]),
        "합계": int(totals["합계"]),
        "은행": info["bank"],
        "계좌번호": info["account"],
        "예금주": info["holder"],
        "출금회사": info["payer"],
        "입금요청일": f"{info['pay_date']:%Y-%m-%d}",
        "긴급": "O" if info["urgent"] else "X",
        "지출결의서 발급": "",
        "완료기안 발급": "",
    }


def record_issue(record: dict, kind: str | None = None) -> bool:
    """같은 건(발급ID)이면 갱신, 없으면 추가. 성공하면 True.
    kind = '지출결의서' / '완료기안' (다운로드 버튼 콜백 — 발급 일시 기록) / None ([저장 및 발급대장 등록])."""
    try:
        df = load_history()
        now = datetime.now().strftime("%Y-%m-%d %H:%M")
        stamp_col = f"{kind} 발급" if kind else None
        hit = df.index[df["발급ID"] == record["발급ID"]]
        if len(hit):
            i = hit[0]
            for col, value in record.items():
                if col.endswith(" 발급"):
                    continue
                if col == "문서번호" and not value:  # 문서번호는 발급 대장에서 입력 → 비어 있으면 기존 값 유지
                    continue
                df.at[i, col] = value
            if stamp_col:
                df.at[i, stamp_col] = now
        else:
            row = dict(record)
            if stamp_col:
                row[stamp_col] = now
            df = pd.concat([df, pd.DataFrame([row])], ignore_index=True)
        save_history(df)
        st.toast(f"발급 대장에 저장했습니다 ({kind or '등록'} · {record['건명']})", icon="📒")
        return True
    except OSError as exc:  # 파일이 Excel에서 열려 있는 경우 등
        st.toast(f"발급 대장 저장 실패: {exc} — history.csv가 다른 프로그램에서 열려 있는지 확인하세요.", icon="⚠️")
        return False


def register_record(record: dict) -> None:
    """[저장 및 발급대장 등록] 콜백: history.csv에 등록 → 불러온 임시저장본은 정리 → 발급 대장 탭으로 이동."""
    if not record_issue(record):
        return
    draft_id = st.session_state.get("current_draft_id")
    if draft_id:
        try:
            save_drafts([d for d in load_drafts() if d["id"] != draft_id])
        except OSError:
            pass  # 임시저장본 정리는 실패해도 등록 자체에는 영향 없음
        st.session_state["current_draft_id"] = None
    st.session_state["main_tab"] = TAB_LEDGER


def history_doc_no(issue_id: str) -> str:
    """발급 대장에 입력된 이 건의 문서번호 (없으면 '')."""
    df = load_history()
    hit = df.loc[df["발급ID"] == issue_id, "문서번호"]
    return str(hit.iloc[0]).strip() if len(hit) else ""


# ---------------------------------------------------------------------------
# 임시저장 (drafts.json)
# ---------------------------------------------------------------------------
def load_drafts() -> list[dict]:
    if not DRAFTS_PATH.exists():
        return []
    try:
        data = json.loads(DRAFTS_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    return data if isinstance(data, list) else []


def save_drafts(drafts: list[dict]) -> None:
    tmp = DRAFTS_PATH.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(drafts, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(DRAFTS_PATH)


def save_draft() -> None:
    """[임시저장] 콜백: 현재 입력값을 저장. 불러온 임시저장본을 수정 중이면 그 항목을 덮어씀."""
    ss = st.session_state
    form = {}
    for key in form_defaults():
        value = ss.get(key)
        form[key] = value.isoformat() if isinstance(value, date) else value
    items_df = ss.get("items_df", blank_items())
    items = [{c: (None if isinstance(v, float) and math.isnan(v) else v) for c, v in r.items()}
             for r in items_df[ITEM_COLUMNS].to_dict("records")]
    filled = [r for r in items if _text(r["품목"])]
    subject = _text(filled[0]["품목"]) if filled else "(품목 없음)"
    total = int(sum(_num(v) for v in items_df["합계"]))
    draft = {"id": ss.get("current_draft_id") or datetime.now().strftime("%y%m%d%H%M%S%f"),
             "saved_at": datetime.now().strftime("%Y-%m-%d %H:%M"), "subject": subject, "total": total,
             "item_count": len(filled), "form": form, "items": items}
    try:
        drafts = [d for d in load_drafts() if d["id"] != draft["id"]]
        save_drafts([draft] + drafts)
    except OSError as exc:
        st.toast(f"임시저장 실패: {exc}", icon="⚠️")
        return
    ss["current_draft_id"] = draft["id"]
    st.toast(f"임시저장했습니다 ({subject}) — [{TAB_DRAFTS}] 탭에서 다시 불러올 수 있습니다.", icon="💾")


def _set_items(df: pd.DataFrame) -> None:
    """내역 표 교체 → 표 위젯을 새 key로 다시 그림 (이전 편집 내용이 새 표에 겹치지 않도록)."""
    st.session_state["items_df"] = df
    st.session_state["items_ver"] = st.session_state.get("items_ver", 0) + 1


def load_draft(draft_id: str) -> None:
    """[불러오기] 콜백: 임시저장본을 입력 폼에 채우고 작성 탭으로 이동."""
    ss = st.session_state
    draft = next((d for d in load_drafts() if d["id"] == draft_id), None)
    if draft is None:
        st.toast("임시저장본을 찾을 수 없습니다.", icon="⚠️")
        return
    defaults = form_defaults()
    for key, default in defaults.items():
        value = draft["form"].get(key, default)
        if key in FORM_DATE_KEYS:
            try:
                value = date.fromisoformat(value)
            except (TypeError, ValueError):
                value = default
        ss[key] = value
    bank = ss["f_bank"]
    if bank and bank not in BANKS:  # 수기 입력한 은행 → 선택 목록에 추가
        ss["custom_banks"] = list(dict.fromkeys(ss.get("custom_banks", []) + [bank]))
    vat_mode = VAT_MODES.get(ss["f_vat"], "exclusive")
    items = pd.DataFrame(draft.get("items") or [{}], columns=EDITOR_COLUMNS)
    _set_items(normalize_items(items, vat_mode))
    ss["current_draft_id"] = draft_id
    ss["main_tab"] = TAB_WRITE
    st.toast(f"임시저장본을 불러왔습니다 ({draft['subject']})", icon="📂")


def delete_draft(draft_id: str) -> None:
    try:
        save_drafts([d for d in load_drafts() if d["id"] != draft_id])
    except OSError as exc:
        st.toast(f"삭제 실패: {exc}", icon="⚠️")
        return
    if st.session_state.get("current_draft_id") == draft_id:
        st.session_state["current_draft_id"] = None
    st.toast("임시저장본을 삭제했습니다.", icon="🗑️")


def new_form() -> None:
    """[새로 작성] 콜백: 입력 폼과 내역 표를 비움."""
    for key, value in form_defaults().items():
        st.session_state[key] = value
    _set_items(blank_items())
    st.session_state["current_draft_id"] = None


def render_drafts() -> None:
    st.header("💾 임시저장 목록")
    st.caption(f"[임시저장]한 작성 중 문서입니다. [불러오기]를 누르면 작성 탭에서 이어서 수정할 수 있습니다. 저장 위치: `{DRAFTS_PATH}`")
    drafts = load_drafts()
    if not drafts:
        st.info(f"임시저장된 문서가 없습니다. [{TAB_WRITE}] 탭에서 [💾 임시저장]을 누르면 여기에 쌓입니다.")
        return
    current = st.session_state.get("current_draft_id")
    for d in drafts:
        with st.container(border=True):
            c1, c2, c3 = st.columns([5, 1.2, 1.2], vertical_alignment="center")
            editing = " · ✏️ 작성 중" if d["id"] == current else ""
            c1.markdown(f"**{d['subject']}**{editing}")
            c1.caption(f"저장 {d['saved_at']} · 품목 {d.get('item_count', 0)}건 · 합계 {int(d.get('total', 0)):,}원")
            c2.button("📂 불러오기", key=f"draft_load_{d['id']}", type="primary", on_click=load_draft, args=(d["id"],), width="stretch")
            c3.button("🗑️ 삭제", key=f"draft_del_{d['id']}", on_click=delete_draft, args=(d["id"],), width="stretch")


def build_history_excel(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "발급대장"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    df = df.assign(상태=history_status(df).str[2:])  # 엑셀에는 색 표시(이모지) 없이
    columns = ["상태"] + [c for c in HISTORY_COLUMNS if c != "발급ID"]
    ws.append(columns)
    for rec in df[columns].to_dict("records"):
        ws.append([rec[c] for c in columns])
    for c, name in enumerate(columns, start=1):
        head = ws.cell(row=1, column=c)
        head.font = Font(name="맑은 고딕", size=10, bold=True)
        head.fill = PatternFill("solid", fgColor="D9E1F2")
        head.alignment = Alignment(horizontal="center", vertical="center")
        width = max([len(str(name)) * 2] + [sum(2 if ord(ch) > 127 else 1 for ch in str(v))
                                           for v in df[name].tolist()]) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(max(width, 8), 50)
        for r in range(1, ws.max_row + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if r > 1:
                cell.font = Font(name="맑은 고딕", size=10)
                if name in HISTORY_NUMBER_COLUMNS:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center" if len(str(cell.value or "")) < 14 else "left",
                                               vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def history_status(df: pd.DataFrame) -> pd.Series:
    """진행 단계: 승인완료 체크 → 3단계, 문서번호 있음(또는 완료기안 발급) → 2단계, 그 외 → 1단계."""
    approved = df["승인완료"].eq("O")
    submitted = df["문서번호"].str.strip().ne("") | df["완료기안 발급"].str.strip().ne("")
    status = pd.Series(STATUS_DRAFT, index=df.index)
    status[submitted] = STATUS_SUBMITTED
    status[approved] = STATUS_DONE
    return status


def apply_ledger_edits(df: pd.DataFrame, original: pd.DataFrame, edited: pd.DataFrame) -> int:
    """대장 표에서 바뀐 문서번호·승인완료를 history.csv에 저장. 바뀐 건수를 돌려줌."""
    def text(v) -> str:  # 칸을 지우면 None/NaN이 올 수 있음
        return "" if v is None or (isinstance(v, float) and math.isnan(v)) else str(v).strip()

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    changed = 0
    for issue_id in edited.index:
        new_doc, old_doc = text(edited.at[issue_id, "문서번호"]), text(original.at[issue_id, "문서번호"])
        new_ok, old_ok = bool(edited.at[issue_id, "승인완료"]), bool(original.at[issue_id, "승인완료"])
        if new_doc == old_doc and new_ok == old_ok:
            continue
        i = df.index[df["발급ID"] == issue_id][0]
        df.at[i, "문서번호"] = new_doc
        if new_ok != old_ok:
            df.at[i, "승인완료"] = "O" if new_ok else ""
            df.at[i, "승인일시"] = now if new_ok else ""
        changed += 1
    if changed:
        save_history(df)
    return changed


def render_ledger() -> None:
    st.header("📒 발급 대장")
    st.caption("[저장 및 발급대장 등록]을 누르거나 지출결의서·완료기안 엑셀을 다운로드하면 자동으로 누적됩니다. "
               f"저장 위치: `{HISTORY_PATH}`")
    df = load_history()
    if df.empty:
        st.info(f"아직 발급 내역이 없습니다. [{TAB_WRITE}] 탭에서 [📒 저장 및 발급대장 등록]을 누르면 여기에 쌓입니다.")
        return
    status = history_status(df)

    # ---- 단계별 요약: 1·2단계는 진행 중인 건명까지, 3단계는 건수만 ----
    for col, stage_label in zip(st.columns(3), (STATUS_DRAFT, STATUS_SUBMITTED, STATUS_DONE)):
        in_stage = df[status == stage_label]
        with col.container(border=True):
            st.metric(stage_label, f"{len(in_stage)}건")
            if stage_label != STATUS_DONE and len(in_stage):
                names = ", ".join(with_gun(" ".join(s.split())) for s in in_stage["건명"] if s.strip())
                st.caption(f"({names})")

    # ---- 필터 ----
    stage = st.radio("진행 단계", list(STATUS_FILTERS), horizontal=True, key="ledger_stage")
    dates = pd.to_datetime(df["작성일"], errors="coerce")
    lo, hi = dates.min().date(), dates.max().date()
    c1, c2, c3 = st.columns([2, 1.5, 1.5])
    period = c1.date_input("작성일 범위", value=(lo, hi), key="ledger_period")
    vendor_q = c2.text_input("지급처 검색", key="ledger_vendor", placeholder="예: 타라")
    subject_q = c3.text_input("건명 검색", key="ledger_subject", placeholder="예: 명함")

    mask = pd.Series(True, index=df.index)
    if STATUS_FILTERS[stage]:
        mask &= status.eq(STATUS_FILTERS[stage])
    if isinstance(period, (tuple, list)) and len(period) == 2:  # 날짜를 하나만 고른 중간 상태는 무시
        mask &= dates.dt.date.between(period[0], period[1])
    if vendor_q.strip():
        mask &= df["지급처"].str.contains(vendor_q.strip(), case=False, regex=False)
    if subject_q.strip():
        mask &= df["건명"].str.contains(subject_q.strip(), case=False, regex=False)

    # 상태 열은 색 동그라미만 (🟡 / 🔵 / 🟢)
    view = df[mask].assign(상태=status[mask].str[0], 승인완료=df.loc[mask, "승인완료"].eq("O"))
    view = view.sort_values(["작성일", "지출결의서 발급"], ascending=False).set_index("발급ID")
    view = view[LEDGER_VIEW_COLUMNS]

    st.caption(f"검색 결과 **{len(view)}건** · 합계 **{int(view['합계'].sum()):,}원** (전체 {len(df)}건) — "
               "표에서 **문서번호**를 입력하거나 **승인완료**를 체크하면 바로 저장됩니다.")
    st.caption(STATUS_LEGEND)
    # 저장 후에는 새 key로 표를 다시 그려, 필터·정렬이 바뀐 뒤 이전 편집 내용이 다른 행에 붙지 않게 함
    version = st.session_state.setdefault("ledger_editor_version", 0)
    edited = st.data_editor(
        view,
        key=f"ledger_editor_{version}",
        hide_index=True,
        width="stretch",
        disabled=[c for c in LEDGER_VIEW_COLUMNS if c not in ("문서번호", "승인완료")],
        column_config={
            "상태": st.column_config.TextColumn("상태", width="small", help=STATUS_LEGEND),
            "건명": st.column_config.TextColumn("건명", width="medium"),
            "승인완료": st.column_config.CheckboxColumn("승인완료", help="회장님 최종 결재·승인이 끝나면 체크"),
            "문서번호": st.column_config.TextColumn("문서번호", help="그룹웨어 상신 후 받은 문서번호",
                                                   width="medium"),
            **{col: st.column_config.NumberColumn(col, format="localized") for col in HISTORY_NUMBER_COLUMNS},
        },
    )
    try:
        changed = apply_ledger_edits(df, view, edited)
    except OSError as exc:
        st.error(f"저장하지 못했습니다: {exc} — history.csv가 Excel 등에서 열려 있으면 닫고 다시 시도하세요.")
        changed = 0
    if changed:
        st.session_state["ledger_editor_version"] = version + 1
        st.toast(f"{changed}건의 진행 상태를 저장했습니다.", icon="✅")
        st.rerun()

    st.download_button(
        "⬇️ 전체 발급 대장 엑셀 다운로드",
        data=build_history_excel(df.sort_values("작성일")),
        file_name=f"발급대장_{date.today():%y%m%d}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        width="stretch",
        key="dl_ledger",
        on_click="ignore",
    )

    with st.expander("🗑️ 발급 내역 삭제 (잘못 저장된 건 정리)"):
        labels = {r["발급ID"]: f"{r['작성일']} · {r['건명']} · {r['지급처']} · {int(r['합계']):,}원"
                  for r in df.to_dict("records")}
        to_delete = st.multiselect("삭제할 내역", list(labels), format_func=labels.get, key="ledger_delete")
        if st.button("선택 내역 삭제", disabled=not to_delete, key="ledger_delete_btn"):
            try:
                make_backup("before-delete")
            except OSError:
                pass  # 백업이 안 돼도 삭제는 진행 (하루 1회 자동 백업이 있음)
            save_history(df[~df["발급ID"].isin(to_delete)])
            st.session_state["ledger_editor_version"] = version + 1
            st.toast(f"{len(to_delete)}건을 삭제했습니다.", icon="🗑️")
            st.rerun()


# ---------------------------------------------------------------------------
# 백업 (backups/*.zip)
# ---------------------------------------------------------------------------
def backup_sources() -> dict[str, Path]:
    """zip 안의 이름 → 실제 파일 경로."""
    return {HISTORY_PATH.name: HISTORY_PATH, DRAFTS_PATH.name: DRAFTS_PATH}


def backup_bytes() -> bytes:
    """현재 데이터 파일(있는 것만)을 zip으로."""
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, path in backup_sources().items():
            if path.exists():
                zf.write(path, name)
    return buf.getvalue()


def make_backup(kind: str) -> Path | None:
    """backups 폴더에 'YYMMDD_HHMMSS_종류.zip' 저장. 백업할 파일이 없으면 None."""
    if not any(p.exists() for p in backup_sources().values()):
        return None
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    path = BACKUP_DIR / f"{datetime.now():%y%m%d_%H%M%S}_{kind}.zip"
    path.write_bytes(backup_bytes())
    for old in sorted(BACKUP_DIR.glob("*_auto.zip"))[:-BACKUP_KEEP]:
        old.unlink(missing_ok=True)
    return path


def daily_backup() -> None:
    """오늘 자동 백업이 없으면 하나 만듦 (앱을 열 때 호출)."""
    if st.session_state.get("daily_backup_done"):
        return
    try:
        if not list(BACKUP_DIR.glob(f"{date.today():%y%m%d}_*_auto.zip")):
            make_backup("auto")
    except OSError as exc:
        st.toast(f"자동 백업 실패: {exc}", icon="⚠️")
    st.session_state["daily_backup_done"] = True


def list_backups() -> list[Path]:
    return sorted(BACKUP_DIR.glob("*.zip"), reverse=True) if BACKUP_DIR.exists() else []


def backup_summary(data: bytes) -> str:
    """zip 안의 대장 건수·임시저장 건수 요약."""
    parts = []
    with zipfile.ZipFile(BytesIO(data)) as zf:
        names = zf.namelist()
        if HISTORY_PATH.name in names:
            df = pd.read_csv(BytesIO(zf.read(HISTORY_PATH.name)), dtype=str, keep_default_na=False, encoding="utf-8-sig")
            parts.append(f"발급 대장 {len(df)}건")
        if DRAFTS_PATH.name in names:
            drafts = json.loads(zf.read(DRAFTS_PATH.name).decode("utf-8"))
            parts.append(f"임시저장 {len(drafts) if isinstance(drafts, list) else 0}건")
    return " · ".join(parts) or "빈 백업"


def restore_backup(data: bytes) -> None:
    """zip의 데이터 파일로 교체. 교체 전에 현재 상태를 'before-restore'로 백업."""
    sources = backup_sources()
    with zipfile.ZipFile(BytesIO(data)) as zf:
        files = {n: zf.read(n) for n in zf.namelist() if n in sources}
    if not files:
        raise ValueError(f"백업 파일에 {' / '.join(sources)}이(가) 없습니다.")
    if HISTORY_PATH.name in files:  # 읽을 수 있는 파일인지 먼저 확인
        pd.read_csv(BytesIO(files[HISTORY_PATH.name]), dtype=str, encoding="utf-8-sig")
    if DRAFTS_PATH.name in files:
        json.loads(files[DRAFTS_PATH.name].decode("utf-8"))
    make_backup("before-restore")
    for name, content in files.items():
        path = sources[name]
        tmp = path.with_suffix(path.suffix + ".tmp")
        tmp.write_bytes(content)
        tmp.replace(path)
    st.session_state["ledger_editor_version"] = st.session_state.get("ledger_editor_version", 0) + 1


def _backup_label(path: Path) -> str:
    stem = path.stem.split("_", 2)
    try:
        when = datetime.strptime(f"{stem[0]}_{stem[1]}", "%y%m%d_%H%M%S").strftime("%Y-%m-%d %H:%M:%S")
    except (ValueError, IndexError):
        return path.name
    kind = stem[2] if len(stem) > 2 else ""
    return f"{when} · {BACKUP_KINDS.get(kind, kind)}"


def _do_restore(data: bytes, label: str) -> None:
    """[복원] 콜백."""
    try:
        restore_backup(data)
    except (OSError, ValueError, zipfile.BadZipFile, pd.errors.ParserError) as exc:
        st.toast(f"복원 실패: {exc} — history.csv가 Excel 등에서 열려 있으면 닫고 다시 시도하세요.", icon="⚠️")
        return
    st.toast(f"복원했습니다 ({label}). 복원 직전 상태도 백업해 두었습니다.", icon="♻️")


def _manual_backup() -> None:
    """[지금 백업] 콜백."""
    try:
        path = make_backup("manual")
    except OSError as exc:
        st.toast(f"백업 실패: {exc}", icon="⚠️")
        return
    if path:
        st.toast(f"백업했습니다: {path.name}", icon="🗄️")
    else:
        st.toast("백업할 데이터가 아직 없습니다.", icon="ℹ️")


def render_backup() -> None:
    st.header("🗄️ 백업")
    st.caption("발급 대장(`history.csv`)과 임시저장(`drafts.json`)을 zip으로 묶어 보관합니다. "
               "앱을 여는 날마다 자동으로 1회 백업하고, 복원·발급 내역 삭제 직전에도 자동 백업합니다. "
               f"저장 위치: `{BACKUP_DIR}`")

    # ---- 현재 데이터 ----
    for col, (label, path) in zip(st.columns(2), (("발급 대장", HISTORY_PATH), ("임시저장", DRAFTS_PATH))):
        with col.container(border=True):
            if path.exists():
                count = len(load_history()) if path == HISTORY_PATH else len(load_drafts())
                st.metric(f"현재 {label}", f"{count}건")
                st.caption(f"마지막 수정 {datetime.fromtimestamp(path.stat().st_mtime):%Y-%m-%d %H:%M}")
            else:
                st.metric(f"현재 {label}", "없음")

    c1, c2 = st.columns(2)
    c1.button("🗄️ 지금 백업", type="primary", width="stretch", on_click=_manual_backup, key="backup_now")
    c2.download_button("⬇️ 현재 데이터 zip 다운로드", data=backup_bytes(),
                       file_name=f"지출결의_백업_{datetime.now():%y%m%d_%H%M}.zip", mime="application/zip",
                       width="stretch", key="backup_download", on_click="ignore")
    st.caption("PC 고장에 대비해 가끔은 zip을 내려받아 USB·메일·클라우드 등 다른 곳에도 보관하세요.")

    # ---- 백업 목록 ----
    st.subheader("백업 목록")
    backups = list_backups()
    if not backups:
        st.info("아직 백업이 없습니다. [🗄️ 지금 백업]을 누르면 여기에 쌓입니다.")
    for path in backups[:30]:
        label = _backup_label(path)
        try:
            data = path.read_bytes()
            summary = backup_summary(data)
        except (OSError, ValueError, zipfile.BadZipFile, pd.errors.ParserError):
            data, summary = None, "⚠️ 읽을 수 없는 파일"
        with st.container(border=True):
            c1, c2, c3 = st.columns([5, 1.2, 1.2], vertical_alignment="center")
            c1.markdown(f"**{label}**")
            c1.caption(f"{summary} · `{path.name}`")
            if data is None:
                continue
            c2.download_button("⬇️ 받기", data=data, file_name=path.name, mime="application/zip",
                               key=f"backup_dl_{path.name}", width="stretch", on_click="ignore")
            with c3.popover("♻️ 복원", width="stretch"):
                st.warning(f"현재 데이터를 **{label}** 시점({summary})으로 되돌립니다. "
                           "지금 상태는 복원 직전에 자동 백업됩니다.")
                st.button("복원 실행", key=f"backup_restore_{path.name}", type="primary",
                          on_click=_do_restore, args=(data, label))
    if len(backups) > 30:
        st.caption(f"최근 30개만 표시합니다 (전체 {len(backups)}개 — `{BACKUP_DIR}` 폴더에서 확인).")

    # ---- 내려받은 zip으로 복원 ----
    with st.expander("📤 내려받은 zip 파일로 복원"):
        up = st.file_uploader("백업 zip 파일", type="zip", key="backup_upload")
        if up is not None:
            data = up.getvalue()
            try:
                st.caption(f"내용: {backup_summary(data)}")
                ok = True
            except (ValueError, zipfile.BadZipFile, pd.errors.ParserError) as exc:
                st.error(f"백업 파일을 읽을 수 없습니다: {exc}")
                ok = False
            st.button("이 파일로 복원", key="backup_restore_upload", type="primary", disabled=not ok,
                      on_click=_do_restore, args=(data, up.name))


# ---------------------------------------------------------------------------
# UI 헬퍼
# ---------------------------------------------------------------------------
def copy_button(text: str, label: str, key: str) -> None:
    """클릭 시 text를 클립보드에 복사하는 버튼 (iframe 내 JS)."""
    payload = json.dumps(text)
    st.iframe(
        f"""
        <button id="btn-{key}" onmouseover="this.style.background='#7A895F';this.style.borderColor='#7A895F'" onmouseout="this.style.background='#8B9A6E';this.style.borderColor='#8B9A6E'" style="
            width:100%;padding:0.5rem 0.75rem;border-radius:0.5rem;cursor:pointer;
            border:1px solid #8B9A6E;background:#8B9A6E;color:#fff;font-weight:600;
            font-size:0.95rem;font-family:sans-serif;">📋 {html.escape(label)}</button>
        <script>
        const btn = document.getElementById("btn-{key}");
        const text = {payload};
        function fallback() {{
            const ta = document.createElement("textarea");
            ta.value = text; ta.style.position = "fixed"; ta.style.opacity = "0";
            document.body.appendChild(ta); ta.focus(); ta.select();
            let ok = false;
            try {{ ok = document.execCommand("copy"); }} catch (e) {{}}
            document.body.removeChild(ta);
            return ok;
        }}
        function done(ok) {{
            btn.textContent = ok ? "✅ 복사 완료!" : "⚠️ 복사 실패 - 직접 선택해 복사하세요";
            setTimeout(() => btn.textContent = "📋 {html.escape(label)}", 1800);
        }}
        btn.addEventListener("click", () => {{
            if (navigator.clipboard && window.isSecureContext) {{
                navigator.clipboard.writeText(text).then(() => done(true), () => done(fallback()));
            }} else {{
                done(fallback());
            }}
        }});
        </script>
        """,
        height=48,
    )


def synced_text(label: str, generated: str, key: str, height: int | None = None) -> str:
    """입력이 바뀌면 자동 재생성하고, 그 전까지는 사용자가 수정한 내용을 유지.

    height가 있으면 text_area, 없으면 한 줄 text_input.
    """
    gen_key = f"{key}__generated"
    if st.session_state.get(gen_key) != generated:
        st.session_state[gen_key] = generated
        st.session_state[key] = generated
    if height is None:
        return st.text_input(label, key=key)
    return st.text_area(label, key=key, height=height)


# ---------------------------------------------------------------------------
# 앱
# ---------------------------------------------------------------------------
def _template_input(label: str, auto: Path | None, help_text: str, key: str) -> Path | None:
    value = st.text_input(label, str(auto) if auto else "", help=help_text, key=key).strip().strip('"')
    return Path(value) if value and Path(value).is_file() else None


def render_sidebar() -> dict:
    with st.sidebar:
        st.header("⚙️ 기본 설정")
        settings = {
            "company": st.text_input("회사", "GSI").strip(),
            "site": st.text_input("사업장", "본사").strip(),
            "dept": st.text_input("부서", "인사총무팀").strip(),
            "drafter": st.text_input("기안자", "김세희 사원").strip(),
        }
        st.divider()
        st.subheader("문자 보고")
        settings["sms_to"] = st.text_input("받는 분", "장미선과장님").strip()
        settings["sms_from"] = st.text_input("보내는 사람 소개", "GSI(주) 총무팀 사원 김세희").strip()
        st.divider()
        settings["unit"] = st.text_input("수량 단위", "개", help="예: 개, 매, 박스, 식")
        st.divider()
        st.subheader("엑셀 템플릿")
        settings["expense_template"] = _template_input(
            "지출결의서 템플릿", find_expense_template(),
            f"기본값: app.py 폴더의 '{EXP_TEMPLATE_NAME}' (없으면 '*지출결의서*.xlsx')", "tpl_expense")
        settings["done_template"] = _template_input(
            "완료기안 템플릿", find_done_template(),
            f"기본값: app.py 폴더의 '{DONE_TEMPLATE_NAME}' (없으면 '*완료기안*.xlsx' 중 원본)", "tpl_done")
    return settings


def render_items_editor(vat_mode: str) -> pd.DataFrame:
    """내역 통합 표 1개: 입력 칸(품목·업체명·수량·단가·배송비·비고) + 자동 계산 칸(순번·공급가액·부가세·합계).

    편집할 때마다 계산 칸을 다시 채운 표로 교체하고, 표 위젯은 새 key로 다시 그림
    (같은 key로 데이터만 바꾸면 이전 편집 내용이 행 추가·삭제 후 엉뚱한 행에 다시 적용됨).
    """
    ss = st.session_state
    if "items_df" not in ss:
        ss["items_df"] = blank_items()
    current = normalize_items(ss["items_df"], vat_mode)
    if not current.equals(ss["items_df"]):  # 부가세 방식이 바뀐 경우
        _set_items(current)
    money = lambda label: st.column_config.NumberColumn(label, format="localized", disabled=True)  # noqa: E731
    edited = st.data_editor(
        ss["items_df"],
        key=f"items_editor_{ss.setdefault('items_ver', 0)}",
        num_rows="fixed",
        width="stretch",
        hide_index=True,
        column_order=EDITOR_COLUMNS,
        column_config={
            "선택": st.column_config.CheckboxColumn("선택", help="삭제할 행을 체크한 뒤 [🗑️ 선택 행 삭제]", width="small"),
            "순번": st.column_config.NumberColumn("순번", format="%d", disabled=True, width="small"),
            "품목": st.column_config.TextColumn("품목", width="large"),
            "업체명": st.column_config.TextColumn("업체명"),
            "수량": st.column_config.NumberColumn("수량", min_value=0, step=1, format="localized"),
            "단가": st.column_config.NumberColumn("단가", min_value=0, step=1, format="localized"),
            "공급가액": money("공급가액"),
            "부가세": money("부가세"),
            "배송비": st.column_config.NumberColumn("배송비", min_value=0, step=1, format="localized"),
            "합계": money("합계"),
            "비고": st.column_config.TextColumn("비고"),
        },
    )
    updated = normalize_items(edited, vat_mode)
    if not updated.equals(ss["items_df"]):
        _set_items(updated)
        st.rerun()

    selected = int(ss["items_df"]["선택"].sum())
    c1, c2, _ = st.columns([1, 1, 3])
    c1.button("➕ 행 추가", key="items_add", width="stretch",
              on_click=lambda: _set_items(pd.concat([ss["items_df"], blank_items()], ignore_index=True)))
    c2.button(f"🗑️ 선택 행 삭제 ({selected})", key="items_delete", width="stretch", disabled=not selected,
              on_click=lambda: _set_items(normalize_items(
                  ss["items_df"][~ss["items_df"]["선택"]].pipe(lambda d: d if len(d) else blank_items()), vat_mode)))
    return ss["items_df"]


def render_writer(settings: dict) -> None:
    # ---- 기능 1: 입력 ----
    ss = st.session_state
    for key, value in form_defaults().items():  # 처음 한 번만 빈 값으로 시작 (샘플 데이터 없음)
        ss.setdefault(key, value)
    st.header("1️⃣ 지출결의서 입력")
    if ss.get("current_draft_id"):
        st.caption("✏️ 임시저장본을 불러와 수정 중입니다. [💾 임시저장]을 누르면 같은 항목에 덮어씁니다.")

    c1, c2, c3, c4 = st.columns(4)
    c1.date_input("작성일", key="f_write_date")
    c2.date_input("입금요청일", key="f_pay_date")
    c3.selectbox("출금회사", ["GSI", "KN"], key="f_payer")
    c4.checkbox("긴급건", key="f_urgent")

    c1, c2, c3 = st.columns(3)
    c1.selectbox("결제방법", PAY_METHODS, key="f_pay_method")
    c2.selectbox("증빙구분", EVIDENCE_TYPES, key="f_evidence", help="'-' = 선택 안 함")
    c3.selectbox("부가세", list(VAT_MODES), key="f_vat")

    c1, c2, c3 = st.columns(3)
    c1.selectbox("은행", BANKS + [b for b in ss.get("custom_banks", []) if b not in BANKS], key="f_bank",
                 index=None, accept_new_options=True, placeholder="선택 또는 직접 입력")
    c2.text_input("계좌번호", key="f_account", placeholder="예: 123-456-7890")
    c3.text_input("예금주", key="f_holder")

    st.subheader("내역")
    st.caption("품목·업체명·수량·단가·배송비·비고를 입력하면 순번·공급가액·부가세·합계가 자동 계산됩니다. "
               "공급가액 = 수량 × 단가 + 배송비 (배송비도 부가세 과세 — 기존 지출결의서·완료기안 양식 기준)")
    vat_mode = VAT_MODES[ss["f_vat"]]
    items_df = render_items_editor(vat_mode)

    reason = st.text_area("사유", key="f_reason", height=90, placeholder="예: 업무 수행에 필요한 물품 구매")
    attachment = st.text_area("첨부", key="f_attachment", height=90, placeholder="예: 가. 견적서 1부.\n나. 통장사본 1부.")
    remark = st.text_area("특이사항", key="f_remark", height=90, placeholder="없으면 비워 두세요")

    items = compute_items(items_df, vat_mode)
    totals = {k: int(items[k].sum()) if not items.empty else 0 for k in ["공급가액", "부가세", "배송비", "합계"]}
    write_date, pay_date, drafter = ss["f_write_date"], ss["f_pay_date"], settings["drafter"]
    info = {
        "write_date": write_date, "drafter": drafter, "doc_no": "", "payer": ss["f_payer"],
        "requester": drafter, "pay_method": ss["f_pay_method"], "evidence": ss["f_evidence"],
        "bank": _text(ss["f_bank"]), "account": ss["f_account"].strip(), "holder": ss["f_holder"].strip(),
        "pay_date": pay_date, "urgent": ss["f_urgent"], "vat_mode": vat_mode, "unit": settings["unit"],
        "subject": item_summary(items),
        "reason": reason.strip(), "attachment": attachment.strip(), "remark": remark.strip(),
        "company": settings["company"], "site": settings["site"], "dept": settings["dept"],
        "sms_to": settings["sms_to"], "sms_from": settings["sms_from"],
    }
    record = make_history_record(info, items, totals) if not items.empty else None

    # ---- 임시저장 / 발급대장 등록 ----
    c1, c2, c3 = st.columns(3)
    c1.button("💾 임시저장", key="btn_draft", width="stretch", on_click=save_draft)
    c2.button("📒 저장 및 발급대장 등록", key="btn_register", type="primary", width="stretch",
              disabled=record is None, on_click=register_record, args=(record,),
              help="history.csv에 등록하고 [2. 발급 대장] 탭으로 이동합니다.")
    c3.button("🆕 새로 작성", key="btn_new", width="stretch", on_click=new_form,
              help="입력 내용을 모두 비웁니다 (임시저장본은 그대로 남음).")

    if items.empty:
        st.info("내역에 품목을 1개 이상 입력하세요.")
        return

    # 문서번호는 발급 대장에서 입력 → 같은 건이 대장에 있으면 그 문서번호를 완료기안에 사용
    doc_no = history_doc_no(record["발급ID"])
    info["doc_no"] = record["문서번호"] = doc_no
    xlsx_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    st.divider()
    # 결과 영역 최상단: 지출결의서 엑셀 (품의 문구·제목 수정본을 반영해야 하므로 아래에서 채움)
    expense_slot = st.container()
    st.divider()
    left, right = st.columns(2)

    # ---- 기능 2: 품의 문구 ----
    with left:
        st.header("2️⃣ 품의서 하단 문구")
        body = synced_text("내용 (수정 가능)", build_body_text(info, items, totals), "body_text", 360)
        copy_button(body, "클립보드 복사", "body")

    # ---- 기능 3 · 5: 제목 / 문자 ----
    with right:
        st.header("3️⃣ 문서 제목")
        title = synced_text("제목 (수정 가능)", build_title(info), "title_text")
        copy_button(title, "제목 복사", "title")

        st.header("5️⃣ 문자 보고 템플릿")
        sms = synced_text("문자 (수정 가능)", build_sms(info, items), "sms_text", 230)
        copy_button(sms, "문자 복사", "sms")

    # ---- 지출결의서 엑셀 (최상단 슬롯) ----
    with expense_slot:
        st.header("📄 지출결의서 엑셀")
        st.caption("원본 템플릿에 값만 채워 넣습니다. 품의 문구(1~4)와 문서 제목은 아래 입력칸에서 수정한 내용이 그대로 들어갑니다. "
                   "다운로드하면 발급 대장에 자동 기록됩니다.")
        expense_template = settings["expense_template"]
        expense_name = f"{write_date:%y%m%d}_지출결의서.xlsx"
        max_items = EXP_ITEM_LAST - EXP_ITEM_FIRST + 1
        template_error = check_expense_template(expense_template) if expense_template else (
            f"app.py 폴더에 지출결의서 템플릿('{EXP_TEMPLATE_NAME}' 또는 '*지출결의서*.xlsx')이 없습니다. "
            "사이드바에서 경로를 지정하세요.")
        if template_error:
            st.error(template_error)
        elif len(items) > max_items:
            st.error(f"양식의 내역 칸은 최대 {max_items}행입니다 (현재 {len(items)}건). 결의서를 나눠 작성하세요.")
        else:
            st.download_button(
                f"⬇️ {expense_name} 다운로드",
                data=build_expense_excel(expense_template, info, items, totals, body, title),
                file_name=expense_name,
                mime=xlsx_mime,
                type="primary",
                width="stretch",
                key="dl_expense",
                on_click=record_issue,
                args=(record, "지출결의서"),
            )
            st.caption(f"템플릿: `{expense_template.name}`")

        st.subheader("미리보기")
        st.caption("[HTML 복사하기] → 그룹웨어 게시판 편집기에 붙여넣으면 표 서식 그대로 들어갑니다. "
                   "(HTML 소스 편집 모드에 붙여넣어도 코드가 그대로 들어감)")
        body_lines = sum(t.count("\n") + 2 for _, t in parse_body_sections(body))
        html_preview(build_expense_html(info, items, totals, body, title),
                     height=700 + 40 * len(items) + 26 * body_lines)

    # ---- 기능 4: 완료기안 엑셀 ----
    st.divider()
    st.header("4️⃣ 완료기안 엑셀")
    st.caption("원본 완료기안 파일의 [입금요청] 시트에 값만 채워 넣습니다. 다운로드하면 발급 대장의 같은 건에 문서번호와 함께 기록됩니다.")
    done_template = settings["done_template"]
    done_name = f"{write_date:%y%m%d}_완료기안 결제요청.xlsx"
    if not done_template:
        st.error(f"app.py 폴더에 완료기안 템플릿('{DONE_TEMPLATE_NAME}' 또는 '*완료기안*.xlsx')이 없습니다. "
                 "사이드바에서 경로를 지정하세요.")
        return
    try:
        done_bytes = load_done_template(str(done_template), done_template.stat().st_mtime)
    except Exception as exc:  # noqa: BLE001 — 템플릿 문제를 화면에 그대로 안내
        st.error(f"완료기안 템플릿 오류: {exc}")
        return
    if not doc_no:
        st.warning(f"문서번호가 비어 있습니다. 상신 후 [{TAB_LEDGER}] 탭에서 이 건의 문서번호를 입력하면 엑셀에 반영됩니다.")
    st.download_button(
        f"⬇️ {done_name} 다운로드",
        data=build_excel(done_bytes, info, items, totals),
        file_name=done_name,
        mime=xlsx_mime,
        type="primary",
        width="stretch",
        key="dl_done",
        on_click=record_issue,
        args=(record, "완료기안"),
    )
    st.caption(f"템플릿: `{done_template.name}` → [{DONE_SHEET}] 시트")


# 브랜드 컬러: 올리브 그린 포인트 + 라이트 그레이 보조
THEME_GREEN, THEME_GREEN_DARK, THEME_BEIGE = "#8B9A6E", "#7A895F", "#EEEEEE"  # Primary / 호버 / 보조 배경
THEME_BORDER, THEME_GREEN_BG, THEME_GREEN_TEXT = "#E0E0E0", "#F2F4EE", "#4A5638"  # 테두리 / 안내 상자 배경·글자
THEME_GREEN_INK = "#5E6B47"  # 흰 배경 위 글자용 (Primary는 글자로 쓰기엔 연해서 한 톤 진하게)


def apply_theme_css() -> None:
    """올리브 그린/라이트 그레이 톤앤매너 CSS — .streamlit/config.toml을 못 읽는 실행 위치에서도 동일하게 적용."""
    st.html(f"""
    <style>
    /* Primary 버튼 (일반·다운로드·폼 제출) */
    button[data-testid="stBaseButton-primary"],
    button[data-testid="stBaseButton-primaryFormSubmit"] {{
        background-color: {THEME_GREEN} !important; border-color: {THEME_GREEN} !important; color: #fff !important;
    }}
    button[data-testid="stBaseButton-primary"] p,
    button[data-testid="stBaseButton-primaryFormSubmit"] p {{ color: #fff !important; font-weight: 600; }}
    button[data-testid="stBaseButton-primary"]:hover:not(:disabled),
    button[data-testid="stBaseButton-primaryFormSubmit"]:hover:not(:disabled),
    button[data-testid="stBaseButton-primary"]:active:not(:disabled) {{
        background-color: {THEME_GREEN_DARK} !important; border-color: {THEME_GREEN_DARK} !important; color: #fff !important;
    }}
    button[data-testid="stBaseButton-primary"]:focus:not(:active) {{
        border-color: {THEME_GREEN_DARK}; box-shadow: 0 0 0 0.2rem {THEME_GREEN}55;
    }}
    /* Secondary 버튼: 호버·포커스 시 올리브 */
    button[data-testid="stBaseButton-secondary"]:hover:not(:disabled),
    button[data-testid="stBaseButton-secondary"]:focus:not(:active) {{
        border-color: {THEME_GREEN}; color: {THEME_GREEN_INK};
    }}
    button[data-testid="stBaseButton-secondary"]:hover:not(:disabled) p {{ color: {THEME_GREEN_INK}; }}
    /* 탭 선택 표시 (글자 + 밑줄) */
    [data-testid="stTab"][aria-selected="true"],
    [data-testid="stTab"][aria-selected="true"] p {{ color: {THEME_GREEN_INK}; font-weight: 600; }}
    [data-testid="stTab"] .react-aria-SelectionIndicator,
    [data-baseweb="tab-highlight"] {{ background-color: {THEME_GREEN} !important; }}
    /* 멀티셀렉트 태그·진행 바·토글 등 기본 빨간 강조 요소 */
    [data-testid="stMultiSelect"] [data-baseweb="tag"] {{ background-color: {THEME_GREEN} !important; color: #fff !important; }}
    [data-testid="stProgress"] [role="progressbar"] > div > div {{ background-color: {THEME_GREEN} !important; }}
    [data-testid="stCheckbox"] label[data-baseweb="checkbox"] input:checked + div {{ background-color: {THEME_GREEN} !important; }}
    /* 입력창·텍스트영역·선택상자·날짜: 라이트 그레이 배경 + 테두리, 포커스 시 올리브 */
    [data-testid="stTextInputRootElement"], [data-testid="stTextAreaRootElement"],
    [data-testid="stNumberInputContainer"], [data-testid="stSelectbox"] [role="group"],
    [data-testid="stMultiSelect"] [role="group"], [data-testid="stDateInputField"] {{
        background-color: {THEME_BEIGE}; border-color: {THEME_BORDER};
    }}
    /* 사이드바는 배경이 라이트 그레이라 입력창을 흰색으로 */
    [data-testid="stSidebar"] [data-testid="stTextInputRootElement"],
    [data-testid="stSidebar"] [data-testid="stSelectbox"] [role="group"] {{
        background-color: #fff;
    }}
    [data-testid="stTextInputRootElement"]:focus-within, [data-testid="stTextAreaRootElement"]:focus-within,
    [data-testid="stNumberInputContainer"]:focus-within, [data-testid="stSelectbox"] [role="group"]:focus-within,
    [data-testid="stMultiSelect"] [role="group"]:focus-within, [data-testid="stDateInputField"]:focus-within {{
        border-color: {THEME_GREEN};
    }}
    /* 안내(info)·성공(success) 박스 */
    [data-testid="stAlert"]:has([data-testid="stAlertContentInfo"]) > div,
    [data-testid="stAlert"]:has([data-testid="stAlertContentSuccess"]) > div {{
        background-color: {THEME_GREEN_BG}; color: {THEME_GREEN_TEXT};
    }}
    [data-testid="stAlertContentInfo"], [data-testid="stAlertContentInfo"] p,
    [data-testid="stAlertContentSuccess"], [data-testid="stAlertContentSuccess"] p {{
        color: {THEME_GREEN_TEXT};
    }}
    /* 테두리 카드(요약·목록): 옅은 그레이 테두리 */
    [data-testid="stVerticalBlockBorderWrapper"], [data-testid="stExpander"] details {{
        border-color: {THEME_BORDER} !important;
    }}
    /* 라디오·체크박스 선택 표시 */
    [data-testid="stRadioOption"][data-selected] > div > div:first-child,
    [data-testid="stCheckbox"] label[data-selected] > span + div {{
        background-color: {THEME_GREEN}; border-color: {THEME_GREEN};
    }}
    </style>
    """)


def main() -> None:
    st.set_page_config(page_title="지출결의 업무", page_icon="🧾", layout="wide")
    apply_theme_css()
    st.title("🧾 지출결의 업무")
    st.caption("지출결의서 내역을 입력하면 품의 문구 · 문서 제목 · 지출결의서/완료기안 엑셀 · 문자 보고 템플릿을 자동 생성하고, "
               "발급 내역을 대장으로 관리합니다.")
    daily_backup()
    settings = render_sidebar()
    # key로 선택 탭을 기억 → [저장 및 발급대장 등록] / [불러오기] 콜백이 session_state로 탭을 바꿈
    tab_write, tab_ledger, tab_drafts, tab_backup = st.tabs([TAB_WRITE, TAB_LEDGER, TAB_DRAFTS, TAB_BACKUP],
                                                            key="main_tab", on_change="rerun")
    with tab_write:
        render_writer(settings)
    with tab_ledger:
        render_ledger()
    with tab_drafts:
        render_drafts()
    with tab_backup:
        render_backup()


if __name__ == "__main__":
    main()
